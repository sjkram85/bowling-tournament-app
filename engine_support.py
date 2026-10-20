import os
import io
import math
import json
import random
import hashlib
import smtplib
import sqlite3
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import pandas as pd
import streamlit as st
import qrcode
from PIL import Image
from pypdf import PdfReader

# Argon2 Cryptographic Hashing
from argon2 import PasswordHasher
from argon2.exceptions import VerifyMismatchError, InvalidHashError

# Word document generation
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image as RLImage, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from google import genai
from google.genai import types

# --- ASSETS & PATHS ---
LOGO_CHUSBC_PATH = "CHUSBC.jpg"
HAS_CHUSBC_LOGO = os.path.exists(LOGO_CHUSBC_PATH)

LOGO_FOREST_HILL_PATH = "forest hill lanes.png"
HAS_FOREST_HILL_LOGO = os.path.exists(LOGO_FOREST_HILL_PATH)

RULEBOOK_PATH = "USBC-Playing-Rulebook-26-27.pdf"
HAS_RULEBOOK = os.path.exists(RULEBOOK_PATH)

PRIMARY_MODEL = "gemini-3.8-flash"
FALLBACK_MODELS = ["gemini-3.7-flash", "gemini-3.5-flash"]
KIOSK_FILE_PATH = os.path.join("/tmp", "chusbc_kiosk_state.json")

LOCAL_DB_PATH = "prism_live.db"

# --- ARGON2 INITIALIZER ---
ph = PasswordHasher(time_cost=2, memory_cost=65536, parallelism=2, hash_len=32, salt_len=16)

def hash_secret(secret_str: str) -> str:
    return ph.hash(str(secret_str).strip())

def verify_secret(stored_hash: str, plain_secret: str) -> tuple[bool, bool]:
    if not stored_hash or not plain_secret:
        return False, False
    clean_secret = str(plain_secret).strip()
    stored_hash = str(stored_hash).strip()
    if stored_hash.startswith("$argon2"):
        try:
            ph.verify(stored_hash, clean_secret)
            return True, ph.check_needs_rehash(stored_hash)
        except (VerifyMismatchError, InvalidHashError):
            return False, False
    else:
        legacy_sha = hashlib.sha256(clean_secret.encode("utf-8")).hexdigest()
        if stored_hash == legacy_sha:
            return True, True
        return False, False

# --- SQLITE LOCAL OPERATIONAL ENGINE ---
def get_sqlite_conn():
    conn = sqlite3.connect(LOCAL_DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn

def init_sqlite_tables():
    conn = get_sqlite_conn()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS live_roster (
            usbc_id TEXT PRIMARY KEY,
            bowler TEXT NOT NULL,
            average INTEGER DEFAULT 200,
            starting_lane INTEGER DEFAULT 1,
            division TEXT DEFAULT 'Open',
            team_name TEXT DEFAULT 'Singles',
            game_1 INTEGER DEFAULT 0,
            game_2 INTEGER DEFAULT 0,
            game_3 INTEGER DEFAULT 0,
            game_4 INTEGER DEFAULT 0,
            paid INTEGER DEFAULT 0,
            checked_in INTEGER DEFAULT 0,
            is_sub INTEGER DEFAULT 0,
            subbing_for TEXT DEFAULT NULL,
            g1_verified_by TEXT DEFAULT '',
            g1_verified_at TEXT DEFAULT '',
            g2_verified_by TEXT DEFAULT '',
            g2_verified_at TEXT DEFAULT '',
            g3_verified_by TEXT DEFAULT '',
            g3_verified_at TEXT DEFAULT '',
            last_updated_by TEXT DEFAULT 'Staff',
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    conn.close()

# --- KIOSK FILE SYSTEM HELPERS ---
def read_kiosk_state():
    if os.path.exists(KIOSK_FILE_PATH):
        try:
            with open(KIOSK_FILE_PATH, "r") as f:
                return json.load(f)
        except Exception:
            return None
    return None

def write_kiosk_state(card_payload):
    try:
        with open(KIOSK_FILE_PATH, "w") as f:
            json.dump(card_payload, f)
    except Exception:
        pass

def clear_kiosk_state():
    try:
        if os.path.exists(KIOSK_FILE_PATH):
            os.remove(KIOSK_FILE_PATH)
    except Exception:
        pass

# --- QR CODE GENERATOR ---
def generate_qr_image(url: str) -> io.BytesIO:
    qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=3)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#0f172a", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf

# --- EMAIL NOTIFICATIONS ---
def notify_admin_new_registration(requested_user: str, requested_name: str, requested_email: str, requested_at: str):
    sender_email = st.secrets.get("EMAIL_SENDER", "")
    sender_password = st.secrets.get("EMAIL_APP_PASSWORD", "")
    target_admin = "scottjk85@gmail.com"
    if not sender_email or not sender_password:
        return False
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = f"🔔 Action Required: New TD Registration ({requested_name})"
        msg["From"] = f"CHUSBC Desk Alert <{sender_email}>"
        msg["To"] = target_admin
        app_url = st.secrets.get("APP_URL", "https://chusbc-desk.streamlit.app").rstrip("/")
        html_body = f"""
        <html><body style="font-family: Arial, sans-serif; color: #0f172a; padding: 20px;">
            <h2>New Staff Registration Pending Approval</h2>
            <p><b>Name:</b> {requested_name}<br/><b>Username:</b> {requested_user}<br/><b>Email:</b> {requested_email}</p>
            <a href="{app_url}" style="background-color: #0284c7; color: #fff; padding: 10px 20px; text-decoration: none; border-radius: 6px;">Sign In to Review</a>
        </body></html>
        """
        msg.attach(MIMEText(html_body, "html"))
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, target_admin, msg.as_string())
        return True
    except Exception as e:
        print(f"Email delivery error: {e}")
        return False

def notify_user_account_approved(target_email: str, target_name: str, username: str):
    sender_email = st.secrets.get("EMAIL_SENDER", "")
    sender_password = st.secrets.get("EMAIL_APP_PASSWORD", "")
    if not sender_email or not sender_password or not target_email:
        return False
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = "🎳 CHUSBC Tournament Desk: Account Approved!"
        msg["From"] = f"CHUSBC Desk Alert <{sender_email}>"
        msg["To"] = target_email
        app_url = st.secrets.get("APP_URL", "https://chusbc-desk.streamlit.app").rstrip("/")
        html_body = f"""
        <html><body style="font-family: Arial, sans-serif; color: #0f172a; padding: 20px;">
            <h2>Account Approved</h2>
            <p>Hello {target_name}, your Tournament Director registration for <b>{username}</b> has been approved.</p>
            <a href="{app_url}" style="background-color: #10b981; color: #fff; padding: 10px 20px; text-decoration: none; border-radius: 6px;">Sign In Now</a>
        </body></html>
        """
        msg.attach(MIMEText(html_body, "html"))
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, target_email, msg.as_string())
        return True
    except Exception as e:
        print(f"Approval email error: {e}")
        return False

# --- LANE BALANCER & CALCULATIONS ---
def parse_active_pairs(range_str: str, max_lanes: int = 40):
    lanes = set()
    parts = [p.strip() for p in range_str.split(",") if p.strip()]
    for part in parts:
        if "-" in part:
            sub = part.split("-")
            try:
                start, end = int(sub[0]), int(sub[1])
                for l in range(start, end + 1):
                    if 1 <= l <= max_lanes:
                        lanes.add(l)
            except ValueError:
                continue
        else:
            try:
                l = int(part)
                if 1 <= l <= max_lanes:
                    lanes.add(l)
            except ValueError:
                continue
    sorted_lanes = sorted(list(lanes))
    pairs = []
    i = 0
    while i < len(sorted_lanes):
        l1 = sorted_lanes[i]
        if l1 % 2 != 0 and (l1 + 1) in lanes:
            pairs.append((l1, l1 + 1))
            i += 2
        else:
            partner = l1 + 1 if l1 % 2 != 0 else l1 - 1
            pairs.append((min(l1, partner), max(l1, partner)))
            i += 1
    seen = set()
    unique = []
    for p in pairs:
        if p not in seen:
            seen.add(p)
            unique.append(p)
    return unique

def get_bowler_lane_for_game(starting_lane: int, game_num: int, movement_rule: str, max_lanes: int = 40) -> int:
    if game_num <= 1:
        return starting_lane
    shift = 0
    if "Move Right 1 Pair" in movement_rule:
        shift = 1
    elif "Move Right 2 Pairs" in movement_rule:
        shift = 2
    elif "Move Right 3 Pairs" in movement_rule:
        shift = 3
    elif "Move Left 1 Pair" in movement_rule:
        shift = -1
    lanes_shift = shift * 2 * (game_num - 1)
    return int(((starting_lane - 1 + lanes_shift) % max_lanes) + 1)

def auto_assign_lanes(roster_df: pd.DataFrame, pairs: list, max_per_pair: int = 5, shuffle: bool = True, group_by_team: bool = True):
    if roster_df.empty or not pairs:
        return roster_df, "No roster or lane pairs provided."
    updated = roster_df.copy()
    if "Team_Name" not in updated.columns:
        updated["Team_Name"] = "Singles"
    updated["Team_Name"] = updated["Team_Name"].fillna("Singles").astype(str)

    if group_by_team:
        team_groups = {}
        for idx, row in updated.iterrows():
            t_name = str(row["Team_Name"]).strip()
            unit_key = t_name if t_name and t_name.lower() not in ["singles", "open", "none", ""] else f"Indiv_{idx}"
            if unit_key not in team_groups:
                team_groups[unit_key] = []
            team_groups[unit_key].append(idx)
        units = list(team_groups.values())
    else:
        units = [[idx] for idx in updated.index]

    if shuffle:
        random.shuffle(units)

    pair_buckets = {p: [] for p in pairs}
    for unit in units:
        unit_size = len(unit)
        assigned = False
        sorted_pairs = sorted(pairs, key=lambda p: len(pair_buckets[p]))
        for p in sorted_pairs:
            if len(pair_buckets[p]) + unit_size <= max_per_pair:
                pair_buckets[p].extend(unit)
                assigned = True
                break
        if not assigned:
            return roster_df, "Unable to fit teams evenly across active pairs."

    for p, row_indices in pair_buckets.items():
        lane_odd, lane_even = p
        for order, r_idx in enumerate(row_indices):
            assigned_lane = lane_odd if (order % 2 == 0) else lane_even
            updated.loc[r_idx, "Starting Lane"] = int(assigned_lane)

    updated["Starting Lane"] = pd.to_numeric(updated["Starting Lane"], errors="coerce").fillna(1).astype(int)
    return updated.sort_values(by=["Starting Lane", "Team_Name", "Bowler"]).reset_index(drop=True), None

def calculate_handicap(avg: int, base: int = 220, percentage: float = 0.90) -> int:
    return 0 if avg >= base else math.floor(percentage * (base - avg))

def compile_standings(df: pd.DataFrame, config: dict = None) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    calc_df = df.copy()
    game_cols = [c for c in ["Game 1", "Game 2", "Game 3", "Game 4"] if c in calc_df.columns]
    for col in game_cols + ["Average"]:
        calc_df[col] = pd.to_numeric(calc_df[col], errors="coerce").fillna(0).astype(int)

    base = int(config.get("hdcp_base", 220)) if config else 220
    pct = float(config.get("hdcp_pct", 0.90)) if config else 0.90
    games_count = int(config.get("games_count", 3)) if config else 3

    calc_df["HDCP_Per_Game"] = calc_df["Average"].apply(lambda a: calculate_handicap(a, base=base, percentage=pct))
    calc_df["Total_HDCP"] = calc_df["HDCP_Per_Game"] * games_count
    active_games = ["Game 1", "Game 2", "Game 3"] if games_count <= 3 else ["Game 1", "Game 2", "Game 3", "Game 4"]
    calc_df["Scratch_Total"] = calc_df[[g for g in active_games if g in calc_df.columns]].sum(axis=1)
    calc_df["Grand_Total"] = calc_df["Scratch_Total"] + calc_df["Total_HDCP"]
    calc_df["High_Scratch_Game"] = calc_df[[g for g in active_games if g in calc_df.columns]].max(axis=1)

    calc_df = calc_df.sort_values(by=["Grand_Total", "High_Scratch_Game", "Scratch_Total"], ascending=[False, False, False]).reset_index(drop=True)
    calc_df.index += 1
    calc_df.index.name = "Rank"
    return calc_df

def calculate_financial_summary(roster_df: pd.DataFrame, config: dict, lineage_per_game: float = 3.75, games_count: int = 3) -> dict:
    if roster_df.empty:
        return {"paid_count": 0, "gross_revenue": 0.0, "total_lineage": 0.0, "net_prize_fund": 0.0}
    paid_count = len(roster_df[roster_df["Paid"] == True])
    entry_fee = float(config.get("entry_fee", 50.0))
    gross = paid_count * entry_fee
    lineage = paid_count * (lineage_per_game * games_count)
    expense = paid_count * 5.00
    return {"paid_count": paid_count, "gross_revenue": gross, "total_lineage": lineage, "net_prize_fund": max(0.0, gross - (lineage + expense))}

def generate_smart_scholarship_csv(standings_df: pd.DataFrame, net_prize_fund: float, tournament_name: str) -> str:
    if standings_df.empty or net_prize_fund <= 0:
        return ""
    top_cut = max(1, math.ceil(len(standings_df) * 0.25))
    eligible = standings_df.head(top_cut).copy()
    weights = [0.40, 0.25, 0.15]
    payouts = []
    for i in range(len(eligible)):
        if i < len(weights):
            amt = round(net_prize_fund * weights[i], 2)
        else:
            remainder = net_prize_fund * (1.0 - sum(weights))
            amt = round(remainder / max(1, (len(eligible) - len(weights))), 2)
        payouts.append(amt)
    eligible["SMART_Award_Amount"] = payouts
    eligible["SMART_Account_Number"] = [f"SMART-{random.randint(100000, 999999)}" for _ in range(len(eligible))]
    return pd.DataFrame({
        "USBC ID": eligible.get("USBC ID", "TBD"), "Bowler Name": eligible["Bowler"],
        "SMART Account #": eligible["SMART_Account_Number"], "Award Amount ($)": eligible["SMART_Award_Amount"],
        "Tournament Name": tournament_name, "Date": datetime.now().strftime("%Y-%m-%d")
    }).to_csv(index=False)

def resolve_bracket_match(b1_id: str, b2_id: str, game_col: str, bracket_type: str, roster_df: pd.DataFrame):
    row1 = roster_df[roster_df["USBC ID"].astype(str) == str(b1_id)]
    row2 = roster_df[roster_df["USBC ID"].astype(str) == str(b2_id)]
    if row1.empty or row2.empty:
        return "-", 0, "-", 0, None
    r1, r2 = row1.iloc[0], row2.iloc[0]
    s1 = int(pd.to_numeric(r1.get(game_col, 0), errors="coerce") or 0)
    s2 = int(pd.to_numeric(r2.get(game_col, 0), errors="coerce") or 0)
    if bracket_type == "Handicap":
        h1 = calculate_handicap(int(pd.to_numeric(r1.get("Average", 200), errors="coerce") or 200))
        h2 = calculate_handicap(int(pd.to_numeric(r2.get("Average", 200), errors="coerce") or 200))
        t1, t2 = s1 + h1, s2 + h2
    else:
        t1, t2 = s1, s2
    winner_id = b1_id if t1 > t2 else b2_id if t2 > t1 else (b1_id if int(r1.get("Average", 0)) >= int(r2.get("Average", 0)) else b2_id)
    return r1["Bowler"], t1, r2["Bowler"], t2, winner_id

# --- REPORT & SOP DOCUMENT GENERATORS ---
def get_report_logo_path(config: dict) -> tuple[str, bool]:
    center_name = str(config.get("center", "")).lower()
    if "forest hill" in center_name and HAS_FOREST_HILL_LOGO:
        return LOGO_FOREST_HILL_PATH, True
    if HAS_CHUSBC_LOGO:
        return LOGO_CHUSBC_PATH, True
    return "", False

def export_lane_assignments_pdf(roster_df: pd.DataFrame, config: dict, td_name: str = "Official") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle("RosterTitle", parent=styles["Heading1"], fontSize=14, leading=17, textColor=colors.HexColor("#0f172a"), spaceAfter=2)
    header_text = f"""
    <b>Cecil Harford USBC — Official Lane Assignment & Check-In Roster</b><br/>
    <font size="8" color="#475569">
    Event: {config['name']} | Host: {config['center']} | Division: {config['division']}<br/>
    Generated: {datetime.now().strftime('%m/%d/%Y %H:%M')} | TD on Duty: {td_name}
    </font>
    """
    pdf_logo_file, has_pdf_logo = get_report_logo_path(config)
    if has_pdf_logo:
        logo_img = RLImage(pdf_logo_file, width=50, height=50)
        header_table = Table([[logo_img, Paragraph(header_text, title_style)]], colWidths=[60, 480])
        header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(header_table)
    else:
        story.append(Paragraph("Cecil Harford USBC — Lane Assignment Roster", title_style))

    if not roster_df.empty:
        df_sorted = roster_df.copy()
        df_sorted["Starting Lane"] = pd.to_numeric(df_sorted.get("Starting Lane", 1), errors="coerce").fillna(1).astype(int)
        df_sorted["Pair_Low"] = df_sorted["Starting Lane"].apply(lambda l: l if l % 2 != 0 else l - 1)
        grouped = df_sorted.groupby("Pair_Low")
        for pair_low, group_df in sorted(grouped, key=lambda x: x[0]):
            pair_name = f"Lanes {int(pair_low)} & {int(pair_low + 1)}"
            p_header = ParagraphStyle("PairHeader", parent=styles["Heading2"], fontSize=10, leading=12, textColor=colors.white)
            pair_table_data = [[Paragraph(f"<b>{pair_name}</b>", p_header), "", "", "", "", "", ""]]
            pair_table_data.append(["Lane", "Bowler Name", "Team / Affiliation", "Avg", "Paid", "Check-in", "Scores (G1 / G2 / G3)"])
            for _, row in group_df.sort_values(by=["Starting Lane", "Bowler"]).iterrows():
                paid = "Yes" if bool(row.get("Paid", False)) else "[  ]"
                checked = "Yes" if bool(row.get("Checked_In", False)) else "[  ]"
                pair_table_data.append([str(row["Starting Lane"]), str(row["Bowler"]), str(row.get("Team_Name", "Singles")), str(row["Average"]), paid, checked, "___ / ___ / ___"])
            pt = Table(pair_table_data, colWidths=[32, 130, 110, 32, 38, 48, 150])
            pt.setStyle(TableStyle([
                ("SPAN", (0, 0), (6, 0)),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#0284c7")),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.white),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(pt)
            story.append(Spacer(1, 8))
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def export_standings_pdf(standings_df: pd.DataFrame, config: dict, champion_name: str = "", td_name: str = "Official") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontSize=14, leading=17, textColor=colors.HexColor("#0f172a"), spaceAfter=2)
    champ_str = f" | <b>CHAMPION: {champion_name}</b>" if champion_name else ""
    header_text = f"""
    <b>Cecil Harford USBC — {config['name']} Official Standings</b><br/>
    <font size="8" color="#475569">
    Host: {config['center']} | Division: {config['division']} | Lanes: {config['lanes_range']}{champ_str}<br/>
    TD on Duty: {td_name} | Generated: {datetime.now().strftime('%m/%d/%Y %H:%M')}
    </font>
    """
    pdf_logo_file, has_pdf_logo = get_report_logo_path(config)
    if has_pdf_logo:
        logo_img = RLImage(pdf_logo_file, width=50, height=50)
        header_table = Table([[logo_img, Paragraph(header_text, title_style)]], colWidths=[60, 480])
        header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(header_table)
    else:
        story.append(Paragraph(f"Cecil Harford USBC — {config['name']}", title_style))

    table_data = [["Pos", "Bowler Name", "USBC ID", "Avg", "HDCP", "G1", "G2", "G3", "Scratch", "Total"]]
    for rank, row in standings_df.reset_index().iterrows():
        table_data.append([
            str(row["Rank"]), str(row["Bowler"]), str(row.get("USBC ID", "-")),
            str(row["Average"]), str(row.get("Total_HDCP", 0)), str(row.get("Game 1", 0)),
            str(row.get("Game 2", 0)), str(row.get("Game 3", 0)), str(row.get("Scratch_Total", 0)), str(row["Grand_Total"])
        ])
    table = Table(table_data, colWidths=[24, 140, 52, 32, 38, 38, 38, 38, 52, 58])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0284c7")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("FONTSIZE", (0, 1), (-1, -1), 7.5),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (-1, -1), "CENTER"),
    ]))
    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_sop_docx(config: dict, tpoc_name: str = "Scott Kramer", tpoc_contact: str = "scottjk85@gmail.com | 443.966.4460") -> io.BytesIO:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    PRIMARY = RGBColor(14, 58, 138)
    p_org = doc.add_paragraph()
    r_org = p_org.add_run("CECIL HARFORD UNITED STATES BOWLING CONGRESS")
    r_org.font.bold = True
    r_org.font.color.rgb = RGBColor(2, 132, 199)
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(f"Standard Operating Procedure: {config.get('name', 'Tournament Desk Operations')}")
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY
    doc.add_paragraph(f"Technical POC: {tpoc_name} ({tpoc_contact}) | Host: {config.get('center', 'Forest Hill Bowling Center')}")
    doc.add_heading(level=1).add_run("1. Desk Setup & Operational Protocols").font.color.rgb = PRIMARY
    doc.add_paragraph(
        "T-60: Connect secondary monitor and load ?mode=kiosk. Load ?mode=leaderboard on concourse display.\n"
        "T-45: Open Module 1, ingest roster, and execute Automated Lane Balancer.\n"
        "T-30: Display entrance self-check-in QR code (?mode=self_checkin). Route unpaid entries to desk.\n"
        "T-15: Roving floor TDs claim pairs via mobile scorer (?mode=scorer) using 4-digit PINs.\n"
        "Squad Active: Floor TDs submit scores. Central Desk audits live progress on the Radar Matrix.\n"
        "Final: Finalize event in Module 4 to execute stepladders and export official hard-copy reports."
    )
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def generate_sop_pdf(config: dict, tpoc_name: str = "Scott Kramer", tpoc_contact: str = "scottjk85@gmail.com | 443.966.4460") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle("SOPTitle", parent=styles["Heading1"], fontSize=15, leading=18, textColor=colors.HexColor("#0e3a8a"))
    body_style = ParagraphStyle("SOPBody", parent=styles["Normal"], fontSize=8.5, leading=12, textColor=colors.HexColor("#0f172a"))
    story.append(Paragraph(f"Cecil Harford USBC — Tournament Desk Standard Operating Procedure", title_style))
    story.append(Paragraph(f"Event: {config.get('name')} | Host: {config.get('center')} | TPOC: {tpoc_name} ({tpoc_contact})", styles["Normal"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("<b>Standard Operations Timeline:</b>", styles["Heading2"]))
    story.append(Paragraph(
        "<b>1. T-60 Monitor Setup:</b> Extend laptop display (Win+P). Open ?mode=kiosk on secondary monitor. Open ?mode=leaderboard on TV.<br/>"
        "<b>2. T-45 Lane Balancing:</b> Balance bowlers across active pairs with team grouping enabled. Print PDF check-in sheets.<br/>"
        "<b>3. T-30 Entrance Check-In:</b> Direct pre-paid bowlers to entrance QR. Route payments to desk.<br/>"
        "<b>4. Scoring & Concurrence:</b> Roving officials enter scores on mobile with 4-digit security PIN concurrence.<br/>"
        "<b>5. Financials & Completion:</b> Reconcile entries against prize fund. Compute stepladders and print summary report.",
        body_style
    ))
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

# --- GEMINI OCR FOR OVERHEAD DISPLAY SNAPSHOTS ---
def extract_scores_from_overhead_photo(ai_client, image_bytes: bytes, bowlers_on_pair: list, game_num: int) -> dict:
    if not ai_client:
        return {}
    prompt = f"""
Analyze this overhead bowling scoring monitor photo.
Target bowlers on this pair: {", ".join(bowlers_on_pair)}
Extract final scores for Game {game_num} (between 0 and 300).
Output strictly raw JSON mapping bowler names to scores: {{"Name": 205}}
"""
    try:
        content_part = types.Part.from_bytes(data=image_bytes, mime_type="image/jpeg")
        response = ai_client.models.generate_content(
            model=PRIMARY_MODEL,
            contents=[content_part, prompt],
            config=types.GenerateContentConfig(temperature=0.0)
        )
        return json.loads(response.text.replace("```json", "").replace("```", "").strip())
    except Exception as e:
        print(f"Monitor OCR Error: {e}")
        return {}

# --- ADD TO BOTTOM OF engine_support.py ---

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image as RLImage, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def export_tournament_summary_pdf(standings_df: pd.DataFrame, config: dict, champion_name: str, stepladder_results: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle("SummaryTitle", parent=styles["Heading1"], fontSize=15, leading=18, textColor=colors.HexColor("#0f172a"), spaceAfter=2)
    h2_style = ParagraphStyle("SummaryH2", parent=styles["Heading2"], fontSize=11, leading=14, textColor=colors.HexColor("#0284c7"), spaceBefore=8, spaceAfter=4)

    header_text = f"""
    <b>Cecil Harford USBC — Official Tournament Summary & Champion Report</b><br/>
    <font size="8" color="#475569">
    Event: {config.get('name', 'CHUSBC Tournament')} | Host Center: {config.get('center', 'Forest Hill')} | Division: {config.get('division', 'Open')}<br/>
    Official Champion: <b>{champion_name}</b> | Generated: {datetime.now().strftime('%m/%d/%Y %H:%M')}
    </font>
    """

    logo_path = LOGO_FOREST_HILL_PATH if ("forest hill" in str(config.get("center", "")).lower() and HAS_FOREST_HILL_LOGO) else LOGO_CHUSBC_PATH
    if os.path.exists(logo_path):
        logo_img = RLImage(logo_path, width=50, height=50)
        header_table = Table([[logo_img, Paragraph(header_text, title_style)]], colWidths=[60, 480])
        header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(header_table)
    else:
        story.append(Paragraph(f"Cecil Harford USBC — {config.get('name', 'Tournament')}", title_style))

    champ_data = [[Paragraph(f"<b>🏆 OFFICIAL CHAMPION CROWNED: {champion_name}</b>", styles["Normal"])]]
    ct = Table(champ_data, colWidths=[540])
    ct.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#10b98120")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#10b981")),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(ct)
    story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def compile_all_events_leaderboard(roster_df: pd.DataFrame, config: dict = None) -> pd.DataFrame:
    if roster_df.empty:
        return pd.DataFrame()
    df = roster_df.copy()
    for col in ["Game 1", "Game 2", "Game 3", "Average"]:
        df[col] = pd.to_numeric(df.get(col, 0), errors="coerce").fillna(0).astype(int)

    base = int(config.get("hdcp_base", 220)) if config else 220
    pct = float(config.get("hdcp_pct", 0.90)) if config else 0.90

    df["HDCP_Per_Game"] = df["Average"].apply(lambda a: calculate_handicap(a, base=base, percentage=pct))
    group_col = "USBC ID" if "USBC ID" in df.columns else "Bowler"
    grouped = df.groupby(group_col).agg({
        "Average": "max", "HDCP_Per_Game": "max", "Game 1": "sum", "Game 2": "sum", "Game 3": "sum",
        "Team_Name": lambda x: ", ".join(set([str(v) for v in x if v and str(v) != "Singles"])) or "Singles",
        "Division": "first"
    }).reset_index()

    if group_col == "USBC ID":
        name_map = df.drop_duplicates(subset=["USBC ID"]).set_index("USBC ID")["Bowler"].to_dict()
        grouped["Bowler"] = grouped["USBC ID"].map(name_map)

    grouped["Total_Games_Bowled"] = df.groupby(group_col).apply(
        lambda g: (g["Game 1"] > 0).sum() + (g["Game 2"] > 0).sum() + (g["Game 3"] > 0).sum()
    ).values
    grouped["Total_HDCP"] = grouped["HDCP_Per_Game"] * grouped["Total_Games_Bowled"]
    grouped["Scratch_Series"] = grouped["Game 1"] + grouped["Game 2"] + grouped["Game 3"]
    grouped["Grand_Total"] = grouped["Scratch_Series"] + grouped["Total_HDCP"]

    grouped = grouped.sort_values(by=["Grand_Total", "Scratch_Series"], ascending=[False, False]).reset_index(drop=True)
    grouped.index += 1
    grouped.index.name = "Rank"
    return grouped
